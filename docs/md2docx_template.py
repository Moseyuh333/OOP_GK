# -*- coding: utf-8 -*-
"""Build SmartSportHub report docx using the thesis template's styles."""
import re, pathlib, sys
sys.path.insert(0, r'C:/Users/MINH KHOI/AppData/Local/hermes/tools')
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = r'C:/Users/MINH KHOI/AppData/Local/hermes/attachments/DeCuongChiTiet_Nhom5_23162026_23162046_NoiDung_v4_CoHinhAnh_APA.docx'
MD = pathlib.Path('D:/New folder/OOP_GK/docs/BAO_CAO_DU_AN.md')
OUT = pathlib.Path('D:/New folder/OOP_GK/docs/BAO_CAO_DU_AN.docx')

doc = Document(TEMPLATE)

# remove all body content, keep sectPr
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, size=12):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'

def add_table(rows):
    if not rows: return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = rows[0]
    tr = t.add_row()
    for i, val in enumerate(hdr):
        set_cell_text(tr.cells[i], val, bold=True, size=12)
    for ridx, row in enumerate(rows[1:]):
        tr = t.add_row()
        for i in range(ncol):
            val = row[i] if i < len(row) else ''
            set_cell_text(tr.cells[i], val, size=12)
    for row in t.rows:
        for c in row.cells:
            c.width = Inches(6.5 / ncol)

def add_code_block(lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    for ln in lines:
        r = p.add_run(ln + '\n')
        r.font.name = 'Consolas'
        r.font.size = Pt(10.5)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'F2F2F2')
    pPr.append(shd)

def inline_runs(p, text):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.size=Pt(12)
        elif part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)

# ---- cover page ----
cover_specs = [
    ('BỘ GIÁO DỤC VÀ ĐÀO TẠO', 13, True),
    ('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH', 13, True),
    ('KHOA CÔNG NGHỆ THÔNG TIN', 13, True),
    ('', 13, False),
    ('', 13, False),
    ('TIỂU LUẬN CHUYÊN NGÀNH', 16, True),
    ('ĐỀ TÀI:', 15, True),
    ('HỆ THỐNG QUẢN LÝ ĐẶT SÂN THỂ THAO — SMARTSPORTHUB', 15, True),
    ('', 13, False),
    ('Sinh viên thực hiện:', 13, True),
    ('Trần Doãn Hòa – MSSV: 23162026', 13, False),
    ('Nguyễn Minh Khôi - MSSV: 23162046', 13, False),
    ('Giảng viên hướng dẫn: Th.S Nguyễn Thị Thanh Vân', 13, False),
    ('', 13, False),
    ('TP. Hồ Chí Minh, tháng 10 năm 2026', 13, False),
]
for t, sz, bold in cover_specs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(t)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(sz)
    run.bold = bold
doc.add_page_break()

# ---- body from markdown ----
lines = MD.read_text(encoding='utf-8').splitlines()
i = 0
in_code = False
code_buf = []
in_table = False
table_buf = []

def flush_table():
    global table_buf
    if table_buf:
        rows = []
        for tb in table_buf:
            cells = [c.strip() for c in tb.strip().strip('|').split('|')]
            if all(re.fullmatch(r'-+', c.replace(' ','')) for c in cells):
                continue
            rows.append(cells)
        add_table(rows)
        table_buf = []

skip_first = True
for line in lines:
    if line.strip() == '# BÁO CÁO ĐỒ ÁN NHÓM':
        continue
    if line.strip().startswith('## HỆ THỐNG QUẢN LÝ ĐẶT SÂN THỂ THAO'):
        continue

    if line.strip().startswith('```'):
        if in_code:
            add_code_block(code_buf); code_buf = []; in_code = False
        else:
            flush_table(); in_table = False
            in_code = True
        continue
    if in_code:
        code_buf.append(line); continue

    if line.strip().startswith('|'):
        in_table = True
        table_buf.append(line); continue
    else:
        if in_table:
            flush_table(); in_table = False

    m = re.match(r'^(#{1,4})\s+(.*)$', line)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        doc.add_heading(text, level=level)
        # fix font color
        for pg in doc.paragraphs[-1].runs:
            pg.font.name = 'Times New Roman'
            pg.font.color.rgb = RGBColor(0x00,0x00,0x00)
        continue

    if line.strip() == '---':
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'000000')
        pbdr.append(bottom); pPr.append(pbdr)
        continue

    # clean TOC markdown links (numbered or bulleted) before list handling
    line = re.sub(r'^(\d+\.\s*|-\s*)\[([^\]]+)\]\([^)]*\)$', r'\1\2', line)

    if re.match(r'^\s*[-*]\s+', line):
        text = re.sub(r'^\s*[-*]\s+', '', line)
        p = doc.add_paragraph(style='List Bullet')
        inline_runs(p, text)
        continue

    if re.match(r'^\s*\d+\.\s+', line):
        text = re.sub(r'^\s*\d+\.\s+', '', line)
        p = doc.add_paragraph(style='List Number')
        inline_runs(p, text)
        continue

    if line.strip() == '':
        continue

    p = doc.add_paragraph()
    inline_runs(p, line)

flush_table()

# add footer page number if not present (template already has)
doc.save(OUT)
print('saved', OUT, OUT.stat().st_size)