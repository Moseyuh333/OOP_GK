# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r'C:/Users/MINH KHOI/AppData/Local/hermes/tools')
from docx import Document
from docx.shared import Pt
p = r'C:/Users/MINH KHOI/AppData/Local/hermes/attachments/DeCuongChiTiet_Nhom5_23162026_23162046_NoiDung_v4_CoHinhAnh_APA.docx'
d = Document(p)
# section
for s in d.sections:
    print('SECTION')
    print(' page width/height:', s.page_width, s.page_height)
    print(' margins L/R/T/B:', s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)
    print(' header:', s.header.paragraphs[0].text if s.header.paragraphs else '')
    print(' footer:', s.footer.paragraphs[0].text if s.footer.paragraphs else '')
# styles
for name in ['Normal','Heading 1','Heading 2','Heading 3','Heading 4','Title','Subtitle','Caption','List Bullet','List Number']:
    st = d.styles[name]
    pf = st.paragraph_format
    f = st.font
    print(f'\nSTYLE {name}:')
    print('  font:', f.name, f.size, 'bold', f.bold, 'italic', f.italic, 'color', f.color.rgb if f.color.type else None)
    print('  space before/after:', pf.space_before, pf.space_after)
    print('  line spacing:', pf.line_spacing)
    print('  left indent:', pf.left_indent, 'first line:', pf.first_line_indent)