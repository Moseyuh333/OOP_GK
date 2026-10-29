# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r'C:/Users/MINH KHOI/AppData/Local/hermes/tools')
from docx import Document
d = Document(r'D:/New folder/OOP_GK/docs/BAO_CAO_DU_AN_APA.docx')
print('paras', len(d.paragraphs), 'tables', len(d.tables))
for p in d.paragraphs[:30]:
    if p.text.strip():
        print(f'[{p.style.name}] {p.text[:70]}')