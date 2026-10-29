# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r'C:/Users/MINH KHOI/AppData/Local/hermes/tools')
from docx import Document
p = r'C:/Users/MINH KHOI/AppData/Local/hermes/attachments/DeCuongChiTiet_Nhom5_23162026_23162046_NoiDung_v4_CoHinhAnh_APA.docx'
d = Document(p)
out = []
for pg in d.paragraphs:
    t = pg.text.strip()
    if t:
        out.append(f'[{pg.style.name}] {t}')
# tables
for ti, tb in enumerate(d.tables):
    out.append(f'--- TABLE {ti} ({len(tb.rows)}x{len(tb.columns)}) ---')
    for row in tb.rows:
        out.append(' | '.join(c.text.strip() for c in row.cells))
open(r'D:/New folder/OOP_GK/docs/template_text.txt','w',encoding='utf-8').write('\n'.join(out))
print('lines', len(out))