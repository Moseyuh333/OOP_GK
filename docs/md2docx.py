# -*- coding: utf-8 -*-
"""Convert docs/BAO_CAO_DU_AN.md to a styled .docx report."""
import re, pathlib
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD = pathlib.Path('D:/New folder/OOP_GK/docs/BAO_CAO_DU_AN.md')
OUT = pathlib.Path('D:/New folder/OOP_GK/docs/BAO_CAO_DU_AN.docx')

doc = Document()

# base style
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

# margins
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, size=10, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    if color: run.font.color.rgb = color

def add_table(rows):
    if not rows: return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # header
    hdr = rows[0]
    tr = t.add_row()
    for i, val in enumerate(hdr):
        set_cell_text(tr.cells[i], val, bold=True, size=10, color=RGBColor(0xFF,0xFF,0xFF))
        shade(tr.cells[i], '1F3864')
    for ridx, row in enumerate(rows[1:]):
        tr = t.add_row()
        for i in range(ncol):
            val = row[i] if i < len(row) else ''
            set_cell_text(tr.cells[i], val, size=10)
            if ridx % 2 == 1:
                shade(tr.cells[i], 'F2F2F2')
    # set column widths roughly
    for row in t.rows:
        for c in row.cells:
            c.width = Inches(6.5 / ncol)

def add_code_block(lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    for ln in lines:
        r = p.add_run(ln + '\n')
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x1E,0x1E,0x1E)
    # shade paragraph
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'F5F5F5')
    pPr.append(shd)

def inline_runs(p, text):
    # handle **bold**, `code`, and normal
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.size=Pt(10)
        elif part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)

lines = MD.read_text(encoding='utf-8').splitlines()
i = 0
in_code = False
code_buf = []
in_table = False
table_buf = []

def flush_table():
    global table_buf
    if table_buf:
        rows = []
        for tb in table_buf:
            cells = [c.strip() for c in tb.strip().strip('|').split('|')]
            if all(re.fullmatch(r'-+', c.replace(' ','')) for c in cells):
                continue
            rows.append(cells)
        add_table(rows)
        table_buf = []

while i < len(lines):
    line = lines[i]
    # code fence
    if line.strip().startswith('```'):
        if in_code:
            add_code_block(code_buf); code_buf = []; in_code = False
        else:
            flush_table(); in_table = False
            in_code = True
        i += 1; continue
    if in_code:
        code_buf.append(line); i += 1; continue
    # table
    if line.strip().startswith('|'):
        flush_table() if False else None
        in_table = True
        table_buf.append(line); i += 1; continue
    else:
        if in_table:
            flush_table(); in_table = False
    # headings
    m = re.match(r'^(#{1,4})\s+(.*)$', line)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        style = {1:'Heading 1',2:'Heading 2',3:'Heading 3',4:'Heading 4'}[level]
        h = doc.add_heading(text, level=level)
        # fix font
        for r in h.runs:
            r.font.name = 'Times New Roman'
            r.font.color.rgb = RGBColor(0x1F,0x38,0x64)
        i += 1; continue
    # horizontal rule
    if line.strip() == '---':
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'1F3864')
        pbdr.append(bottom); pPr.append(pbdr)
        i += 1; continue
    # bullet
    if re.match(r'^\s*[-*]\s+', line):
        text = re.sub(r'^\s*[-*]\s+', '', line)
        p = doc.add_paragraph(style='List Bullet')
        inline_runs(p, text)
        i += 1; continue
    # numbered list
    if re.match(r'^\s*\d+\.\s+', line):
        text = re.sub(r'^\s*\d+\.\s+', '', line)
        p = doc.add_paragraph(style='List Number')
        inline_runs(p, text)
        i += 1; continue
    # blank
    if line.strip() == '':
        i += 1; continue
    # normal paragraph
    p = doc.add_paragraph()
    inline_runs(p, line)
    i += 1

flush_table()
doc.save(OUT)
print('saved', OUT, OUT.stat().st_size, 'bytes')