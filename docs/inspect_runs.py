# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r'C:/Users/MINH KHOI/AppData/Local/hermes/tools')
from docx import Document
p = r'C:/Users/MINH KHOI/AppData/Local/hermes/attachments/DeCuongChiTiet_Nhom5_23162026_23162046_NoiDung_v4_CoHinhAnh_APA.docx'
d = Document(p)
tb = d.tables[0]
for c in tb.rows[0].cells:
    for pg in c.paragraphs:
        for run in pg.runs:
            print('text:', repr(run.text), 'bold:', run.bold, 'size:', run.font.size, 'name:', run.font.name, 'color:', run.font.color.rgb if run.font.color.type else None)
    print('---')