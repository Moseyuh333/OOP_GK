# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r'C:/Users/MINH KHOI/AppData/Local/hermes/tools')
from docx import Document
p = r'C:/Users/MINH KHOI/AppData/Local/hermes/attachments/DeCuongChiTiet_Nhom5_23162026_23162046_NoiDung_v4_CoHinhAnh_APA.docx'
d = Document(p)
print('paragraphs', len(d.paragraphs), 'tables', len(d.tables))
print('---- styles ----')
for s in d.styles:
    try:
        if s.type is not None:
            print(repr(s.name), s.type)
    except Exception as e:
        print('err', e)
print('---- first 60 paragraphs ----')
for pgh in d.paragraphs[:60]:
    t = pgh.text.strip()
    if t:
        print(f'[{pgh.style.name}] {t[:120]}')